"""Build a short advisor-report deck (Walrus base vs TRL2D-fine-tuned, rel L1/L2 focus): 4 slides.

Uses the user's existing deck as the template and reproduces its slide formatting at
the XML level: a 28 pt bold date text box at the top-left, a body text box with a
numbered first level (arabicPeriod, hanging indent) and Wingdings bullets on the
second level, 150 % line spacing, Times New Roman latin font with bold / red
emphasis, and tables in the deck's default table style with Times New Roman text.

Example::

    python scripts/make_report_pptx.py runs/eval_trl2d runs/eval_trl2d_ft \\
        --template D:/lrwei/ppt/ppt.pptx --date 2026.09.02 \\
        --out D:/lrwei/ppt/walrus_trl2d_report.pptx --fig-dir runs/report_figs
"""

from __future__ import annotations

import argparse
import copy
import os
import sys

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Emu, Pt  # noqa: E402

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from export_excel import PAPER, load_curves, load_loss_dicts  # noqa: E402

FIELDS = ["density", "pressure", "velocity_x", "velocity_y"]
FIELD_CN = {"density": "density（密度）", "pressure": "pressure（压力）", "velocity_x": "velocity_x", "velocity_y": "velocity_y", "full": "四字段平均"}
RANGES = [("0:1", "第 1 步"), ("1:3", "第 2–3 步"), ("3:9", "第 4–9 步"), ("9:30", "第 10–30 步"), ("30:95", "第 31–95 步"), ("all", "全程 1–95 步")]
BLUE, ORANGE = "#2a78d6", "#eb6834"
RED = RGBColor(0xFF, 0x00, 0x00)
GOOD = RGBColor(0x15, 0x80, 0x3D)  # green / red used in the reference deck's tables
BAD = RGBColor(0xB9, 0x1C, 0x1C)
TNR = "Times New Roman"

# geometry copied from the reference deck (EMU)
DATE_OFF, DATE_EXT = (748145, 480291), (2807855, 523220)
BODY_OFF = (748145, 891562)
LVL0_PPR = '<a:pPr xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" marL="342900" indent="-342900"><a:lnSpc><a:spcPct val="150000"/></a:lnSpc><a:buFontTx/><a:buAutoNum type="arabicPeriod"/></a:pPr>'
LVL1_PPR = '<a:pPr xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" marL="742950" lvl="1" indent="-285750"><a:lnSpc><a:spcPct val="150000"/></a:lnSpc><a:buFont typeface="Wingdings" panose="05000000000000000000" pitchFamily="2" charset="2"/><a:buChar char="l"/></a:pPr>'
NOTE_PPR = '<a:pPr xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" marL="742950" lvl="1" indent="-285750"><a:buFont typeface="Wingdings" panose="05000000000000000000" pitchFamily="2" charset="2"/><a:buChar char="l"/></a:pPr>'


# ------------------------------------------------------------------ figure
def style_axes(ax, ylabel: str):
    ax.set_yscale("log")
    ax.grid(True, which="major", color="#d9e0e8", linewidth=0.6)
    ax.grid(True, which="minor", color="#eef2f6", linewidth=0.4)
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    for side in ("left", "bottom"):
        ax.spines[side].set_color("#b8c2cc")
        ax.spines[side].set_linewidth(0.6)
    ax.tick_params(colors="#52514e", labelsize=8, width=0.6)
    ax.set_xlabel("rollout step", fontsize=9, color="#52514e")
    ax.set_ylabel(ylabel, fontsize=9, color="#52514e")
    ax.set_xlim(0, 96)


def fig_rollout(curves, labels, out):
    """RelL2 / RelL1 vs step, base vs fine-tuned (test, per-step median over 9 trajectories)."""
    fig, axes = plt.subplots(1, 2, figsize=(7.4, 3.2), dpi=200)
    for ax, met in zip(axes, ["RelL2", "RelL1"]):
        for lab, color in zip(labels, (BLUE, ORANGE)):
            y = curves[lab][("full", met, "median")]
            ax.plot(range(1, len(y) + 1), y, color=color, linewidth=1.2, marker="o", markersize=2.2, label=lab)
        style_axes(ax, met)
        ax.set_title(f"{met} vs rollout step", fontsize=10)
    axes[0].legend(frameon=False, fontsize=9, loc="lower right")
    fig.tight_layout()
    fig.savefig(out)
    plt.close(fig)


# ------------------------------------------------------------------ pptx helpers
def clear_slides(prs):
    sld_ids = prs.slides._sldIdLst
    for sld_id in list(sld_ids):
        prs.part.drop_rel(sld_id.rId)
        sld_ids.remove(sld_id)


def _set_run(run, text, size, bold=False, color=None, latin=TNR):
    run.text = text
    rpr = run._r.get_or_add_rPr()
    rpr.set("lang", "zh-CN")
    rpr.set("altLang", "en-US")
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if latin:
        run.font.name = latin


def new_slide(prs, date: str):
    """Blank slide from the deck's first layout with the reference date box (28 pt bold, theme font)."""
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    for ph in list(slide.placeholders):
        ph._element.getparent().remove(ph._element)
    box = slide.shapes.add_textbox(Emu(DATE_OFF[0]), Emu(DATE_OFF[1]), Emu(DATE_EXT[0]), Emu(DATE_EXT[1]))
    box.text_frame.word_wrap = True
    box.text_frame.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    run = box.text_frame.paragraphs[0].add_run()
    run.text = date
    rpr = run._r.get_or_add_rPr()
    rpr.set("lang", "en-US")
    rpr.set("altLang", "zh-CN")
    run.font.size = Pt(28)
    run.font.bold = True
    return slide


def add_text(slide, lines, top=None, width=840, size=16, note=False):
    """Body text box in the reference format.

    lines: (text, level) with level 0 = numbered item, 1 = Wingdings bullet.  '**' toggles bold red.
    note=True: bullets without the 150 % line spacing (small footnotes under a table / figure).
    """
    top_emu = BODY_OFF[1] if top is None else Pt(top)
    box = slide.shapes.add_textbox(Emu(BODY_OFF[0]), top_emu, Pt(width), Pt(40))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    for k, (text, level) in enumerate(lines):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        ppr_xml = NOTE_PPR if note else (LVL0_PPR if level == 0 else LVL1_PPR)
        p._p.insert(0, etree.fromstring(ppr_xml))
        for i, seg in enumerate(text.split("**")):
            if not seg:
                continue
            emphasised = i % 2 == 1
            _set_run(p.add_run(), seg, size, bold=emphasised, color=RED if emphasised else None)
    return box


BORDER_XML = (
    '<a:{side} xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" w="6350" cap="flat" cmpd="sng" algn="ctr">'
    '<a:solidFill><a:srgbClr val="000000"/></a:solidFill><a:prstDash val="solid"/></a:{side}>'
)


def add_table(slide, header, rows, left, top, col_widths, font_size=11, row_height=None, pct_cols=()):
    """Plain grid like the deck's own tables: white cells, 0.5 pt black borders, Times New Roman, bold header."""
    n_rows, n_cols = len(rows) + 1, len(header)
    rh = row_height or (font_size + 8)
    shape = slide.shapes.add_table(n_rows, n_cols, Pt(left), Pt(top), Pt(sum(col_widths)), Pt(rh * n_rows))
    table = shape.table
    tbl_pr = shape._element.graphic.graphicData.tbl.tblPr
    for flag in ("firstRow", "bandRow"):  # no header shading / banding in the reference tables
        tbl_pr.attrib.pop(flag, None)
    for j, w in enumerate(col_widths):
        table.columns[j].width = Pt(w)
    for r in range(n_rows):
        table.rows[r].height = Pt(rh)
        for c in range(n_cols):
            cell = table.cell(r, c)
            cell.margin_left = cell.margin_right = Emu(34811)
            cell.margin_top = cell.margin_bottom = Emu(17405)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            tc_pr = cell._tc.get_or_add_tcPr()
            for k, side in enumerate(("lnL", "lnR", "lnT", "lnB")):  # borders must precede the fill element
                tc_pr.insert(k, etree.fromstring(BORDER_XML.format(side=side)))
            value = header[c] if r == 0 else rows[r - 1][c]
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if (c == 0 and r > 0) else PP_ALIGN.CENTER
            is_pct = c in pct_cols and r > 0 and isinstance(value, float)
            text = f"{value:+.0%}" if is_pct else (f"{value:.4f}" if isinstance(value, float) else str(value))
            color = (GOOD if value > 0 else BAD) if is_pct else None
            _set_run(p.add_run(), text, font_size, bold=(r == 0 or is_pct), color=color)
    return shape


def imp(base: float, ft: float) -> float:
    return (base - ft) / base


def window_mean(curve, a, b):  # 1-indexed inclusive
    return sum(curve[a - 1 : b]) / (b - a + 1)


# ------------------------------------------------------------------ main
def main() -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("base_run")
    p.add_argument("ft_run")
    p.add_argument("--template", required=True, help="existing deck whose theme / size / style to reuse")
    p.add_argument("--date", required=True, help="date shown at the top-left of every slide, e.g. 2026.09.02")
    p.add_argument("--out", required=True)
    p.add_argument("--fig-dir", default="runs/report_figs")
    args = p.parse_args()
    B, F = "base", "finetuned"
    loss = {B: load_loss_dicts(args.base_run), F: load_loss_dicts(args.ft_run)}
    curves = {B: load_curves(args.base_run, "test"), F: load_curves(args.ft_run, "test")}
    os.makedirs(args.fig_dir, exist_ok=True)
    f_roll = os.path.join(args.fig_dir, "rollout_rel.png")
    fig_rollout(curves, [B, F], f_roll)

    def v(lab, split, field, met, agg="mean", time="all"):
        return loss[lab][split][(field, met, time, agg)]

    rl1_b, rl1_f = v(B, "test", "full", "RelL1"), v(F, "test", "full", "RelL1")
    rl2_b, rl2_f = v(B, "test", "full", "RelL2"), v(F, "test", "full", "RelL2")
    c2 = {lab: curves[lab][("full", "RelL2", "median")] for lab in (B, F)}
    c1 = {lab: curves[lab][("full", "RelL1", "median")] for lab in (B, F)}
    cv = {lab: curves[lab][("full", "VRMSE", "median")] for lab in (B, F)}

    prs = Presentation(args.template)
    clear_slides(prs)

    # ---- 1 结论
    s = new_slide(prs, args.date)
    add_text(s, [
        ("Walrus 在 TRL2D 上的实际相对误差（预训练 base vs 官方微调版）", 0),
        ("数据：The Well 二维湍流辐射层，128×384 网格，density / pressure / velocity 四个场，test 9 条轨迹 × 101 帧；上下文 6 帧预测下一帧；误差在反归一化后的原始物理量上计算", 1),
        (f"一步预测（855 个样本均值）：rel L2 base {rl2_b:.1%} → 微调 **{rl2_f:.1%}**，rel L1 {rl1_b:.1%} → **{rl1_f:.1%}**", 1),
        (f"自回归 rollout（微调版）：前 20 步平均 rel L2 **{window_mean(c2[F], 1, 20):.1%}**，第 21–60 步 **{window_mean(c2[F], 21, 60):.1%}**；30 步后误差饱和，微调只改善短程", 1),
        ("与论文 VRMSE 对表结果一致（第 4 页），评测流程可信", 1),
    ])
    rows = []
    for name, cb, cf, vb, vf in [("rel L2（四场展平）", c2[B], c2[F], rl2_b, rl2_f), ("rel L1（四场展平）", c1[B], c1[F], rl1_b, rl1_f)]:
        rows.append([name, vb, vf, window_mean(cb, 1, 20), window_mean(cf, 1, 20), window_mean(cb, 21, 60), window_mean(cf, 21, 60)])
    add_table(s, ["指标", "一步 base", "一步 微调", "T[1:20] base", "T[1:20] 微调", "T[21:60] base", "T[21:60] 微调"], rows,
              left=59, top=300, col_widths=[180, 110, 110, 110, 110, 110, 110])
    add_text(s, [("一步 = test 全部滑窗样本的均值；T[a:b] = rollout 第 a–b 步（每步先对 9 条轨迹取中位数再平均）", 1)], top=372, size=12, note=True)

    # ---- 2 一步预测逐字段
    s = new_slide(prs, args.date)
    add_text(s, [
        ("一步预测：逐字段与全场展平的相对误差（test，855 个样本均值）", 0),
        (f"density 改善最大（相对 L2 {v(B, 'test', 'density', 'NRMSE'):.1%} → {v(F, 'test', 'density', 'NRMSE'):.1%}），速度场最难（6–8%），pressure 最小（<1%）；展平指标由量级最大的 density 主导", 1),
    ])
    header = ["字段", "相对 L1 base", "相对 L1 微调", "提升率", "相对 L2 base", "相对 L2 微调", "提升率"]
    rows = []
    for f in FIELDS + ["full"]:
        r = [FIELD_CN[f]]
        for met in ("NMAE", "NRMSE"):
            b, ft = v(B, "test", f, met), v(F, "test", f, met)
            r += [b, ft, imp(b, ft)]
        rows.append(r)
    rows.append(["全场展平（RelL1 | RelL2）", rl1_b, rl1_f, imp(rl1_b, rl1_f), rl2_b, rl2_f, imp(rl2_b, rl2_f)])
    add_table(s, header, rows, left=59, top=175, col_widths=[220, 105, 105, 90, 105, 105, 90], pct_cols=(3, 6))
    add_text(s, [
        ("逐字段：相对 L1 = Σ|x−y| / Σ|y|，相对 L2 = ‖x−y‖ / ‖y‖，每个场单独算（The Well 的 NMAE / NRMSE）；全场展平 = 四个场与全部空间点拼成一个向量再算（常用的 rel L1 / rel L2）", 1),
        ("提升率 = (base − 微调) / base，正值为改善", 1),
    ], top=350, size=12, note=True)

    # ---- 3 rollout
    s = new_slide(prs, args.date)
    add_text(s, [
        ("自回归 rollout：rel L1 / L2 随步数变化（test，9 条轨迹，对数纵轴）", 0),
        (f"微调版 rel L2：第 1 步 {c2[F][0]:.1%} → 第 10 步 {c2[F][9]:.1%} → 第 20 步 {c2[F][19]:.1%} → 第 60 步 {c2[F][59]:.1%}；前 10 步误差比 base 减半以上，30 步后两者接近（湍流混合层混沌，逐点误差必然饱和）", 1),
    ])
    s.shapes.add_picture(f_roll, Pt(45), Pt(175), width=Pt(520))
    rows = []
    for key, name in RANGES:
        rows.append([name, v(B, "rollout_test", "full", "RelL2", "mean", key), v(F, "rollout_test", "full", "RelL2", "mean", key),
                     v(B, "rollout_test", "full", "RelL1", "mean", key), v(F, "rollout_test", "full", "RelL1", "mean", key)])
    add_table(s, ["步数区间", "rel L2 base", "rel L2 微调", "rel L1 base", "rel L1 微调"], rows, left=580, top=185, col_widths=[100, 68, 68, 68, 68], font_size=11)
    add_text(s, [("图：每步对 9 条轨迹取中位数。表：区间内先按步平均，再对 9 条轨迹取均值。", 1)], top=405, size=12, note=True)

    # ---- 4 与论文对表 / VRMSE vs rel L2
    s = new_slide(prs, args.date)
    one_b, one_f = v(B, "test", "full", "VRMSE", "median"), v(F, "test", "full", "VRMSE", "median")
    add_text(s, [
        ("与论文对表（VRMSE），以及 VRMSE 为什么比 rel L2 大得多", 0),
        ("同一权重、同一数据、同一指标下，本地 VRMSE 与论文 Tables 13–15 同量级且略优（差异来自聚合顺序和本地只有 9 条 rollout 轨迹）→ 评测流程可信", 1),
        ("VRMSE = RMSE / std(y)，只对波动归一化、对常数偏移敏感：pressure 场几乎均匀（std 0.06），相对 L2 只有 0.6%，VRMSE 却有 0.19；rel L2 才反映真实的逐点误差水平", 1),
    ])
    rows = [
        ["one-step（一步预测，855 样本中位数）", PAPER["1:1"], one_f, one_b, v(F, "test", "full", "RelL2", "median"), v(B, "test", "full", "RelL2", "median")],
        ["T[1:20]", PAPER["1:20"], window_mean(cv[F], 1, 20), window_mean(cv[B], 1, 20), window_mean(c2[F], 1, 20), window_mean(c2[B], 1, 20)],
        ["T[21:60]", PAPER["21:60"], window_mean(cv[F], 21, 60), window_mean(cv[B], 21, 60), window_mean(c2[F], 21, 60), window_mean(c2[B], 21, 60)],
    ]
    add_table(s, ["时间窗口", "论文 VRMSE 微调", "本地 VRMSE 微调", "本地 VRMSE base", "本地 rel L2 微调", "本地 rel L2 base"], rows,
              left=59, top=235, col_widths=[250, 118, 118, 118, 118, 118], font_size=11)
    rows = [[FIELD_CN[f], v(F, "test", f, "NRMSE"), v(F, "test", f, "VRMSE")] for f in FIELDS]
    add_table(s, ["微调版一步预测（均值）", "相对 L2", "VRMSE"], rows, left=59, top=345, col_widths=[250, 118, 118], font_size=11)

    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    prs.save(args.out)
    print(f"wrote {args.out}: {len(prs.slides)} slides; figure {f_roll}")


if __name__ == "__main__":
    main()
